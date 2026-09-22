from __future__ import print_function
import io
import openpyxl
from google.oauth2.service_account import Credentials
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload

SCOPES = [
    "https://www.googleapis.com/auth/spreadsheets.readonly",
    "https://www.googleapis.com/auth/drive.readonly",
]
COLUMN_MAP = {
    "email": "email",
    "full_name": "full_name",
    "category": "category",
    "language": "language",
    "source": "source",
    "notes": "notes",
}
CREDENTIALS_FILE = "app/Credentials/Credentials.json"
GOOGLE_SHEET_MIMETYPE = "application/vnd.google-apps.spreadsheet"
XLSX_MIMETYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def normalize_headers(raw_headers) -> dict:
    mapping = {}
    unrecognized = []
    for i, header in enumerate(raw_headers):
        if header is None:
            continue
        normalized = header.strip().lower()
        if normalized in COLUMN_MAP:
            mapping[i] = COLUMN_MAP[normalized]
        else:
            unrecognized.append(header)

    missing = [col for col in COLUMN_MAP if col not in mapping.values()]
    print(f"Recognized {len(mapping)}/{len(COLUMN_MAP)} expected columns")
    if unrecognized:
        print(f"Unrecognized headers in sheet (ignored): {unrecognized}")
    if missing:
        print(f"Expected columns not found in sheet: {missing}")
    return mapping


def get_drive_services():
    print(f"Auth with service account file: {CREDENTIALS_FILE}")
    creds = Credentials.from_service_account_file(CREDENTIALS_FILE, scopes=SCOPES)
    return build("drive", "v3", credentials=creds, always_use_jwt_access=False)


def find_file(drive_service, file_name: str) -> dict:
    query = f"name='{file_name}' and trashed=false"
    print(f"Searching drive with query: {query}")
    results = drive_service.files().list(
        q=query,
        fields="files(id, name, mimeType, modifiedTime, owners, parents)",
        spaces="drive"
    ).execute()

    files = results.get("files", [])
    if not files:
        print(f"No file named '{file_name}' found (check sharing/permissions on the service account)")
        raise FileNotFoundError(f"No file named '{file_name}' found in Drive")

    if len(files) > 1:
        print(f"{len(files)} files named '{file_name}' found — using the first match. Candidates:")
        for f in files:
            print(f"  id={f['id']} mimeType={f['mimeType']} modified={f.get('modifiedTime')}")
    else:
        print(
            f"Found file: id={files[0]['id']} mimeType={files[0]['mimeType']} "
            f"modified={files[0].get('modifiedTime')}"
        )
    return files[0]


def download_excel_to_memory(drive_service, file_meta: dict) -> io.BytesIO:
    file_id = file_meta["id"]
    mime_type = file_meta["mimeType"]

    if mime_type == GOOGLE_SHEET_MIMETYPE:
        print("File is a native Google Sheet — using export instead of get_media")
        request = drive_service.files().export_media(fileId=file_id, mimeType=XLSX_MIMETYPE)
    elif mime_type == XLSX_MIMETYPE:
        print("File is an uploaded .xlsx — using get_media")
        request = drive_service.files().get_media(fileId=file_id)
    else:
        print(f"Unexpected mimeType '{mime_type}' — cannot read as Excel")
        raise ValueError(f"Unsupported file type: {mime_type}")

    buffer = io.BytesIO()
    downloader = MediaIoBaseDownload(buffer, request)
    done = False
    while not done:
        status, done = downloader.next_chunk()
        if status:
            print(f"Download progress: {int(status.progress() * 100)}%")

    buffer.seek(0)
    print(f"Downloaded {buffer.getbuffer().nbytes} bytes into memory")
    return buffer


def print_emails(data_rows, column_map):
    """Prints just the email column, using column_map to find its index."""
    email_index = None
    for idx, field in column_map.items():
        if field == "email":
            email_index = idx
            break

    if email_index is None:
        print("No 'email' column found in this sheet — check headers/COLUMN_MAP")
        return

    print("\nEmails found:")
    print("-" * 50)
    count = 0
    for row in data_rows:
        if email_index < len(row) and row[email_index]:
            print(row[email_index])
            count += 1
    print("-" * 50)
    print(f"Total emails: {count}")


def import_from_drive(file_name: str = "test_contacts.xlsx"):
    drive_service = get_drive_services()
    file_meta = find_file(drive_service, file_name)
    buffer = download_excel_to_memory(drive_service, file_meta)

    workbook = openpyxl.load_workbook(buffer, read_only=True, data_only=True)
    sheet = workbook.active
    print(f"Reading sheet: '{sheet.title}'")
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        print("Sheet is empty")
        return {"added": 0, "skipped": 0, "invalid": 0}

    headers = rows[0]
    column_map = normalize_headers(list(headers))
    data_rows = rows[1:]

    print(f"Headers: {headers}")
    print(f"Column map (index -> field): {column_map}")
    print(f"Data rows found: {len(data_rows)}")

    print_emails(data_rows, column_map)

    return data_rows, column_map


def print_excel_contents(file_name: str = "test_contacts.xlsx"):
    drive_service = get_drive_services()
    file_meta = find_file(drive_service, file_name)
    buffer = download_excel_to_memory(drive_service, file_meta)

    workbook = openpyxl.load_workbook(buffer, read_only=True, data_only=True)
    sheet = workbook.active
    print(f"\nSheet: {sheet.title}")
    print("-" * 50)

    rows = sheet.iter_rows(values_only=True)
    headers = next(rows, None)
    if headers is None:
        print("Sheet is empty")
        workbook.close()
        return

    print("Headers:")
    print(headers)
    print("\nDATA:")
    print("-" * 50)

    for row_number, row in enumerate(rows, start=1):
        print(f"Row {row_number}: {row}")

    workbook.close()


if __name__ == "__main__":
    import_from_drive("test_contacts.xlsx")
